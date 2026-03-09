from __future__ import annotations

import csv
import html
import json
import re
from dataclasses import dataclass, field
from datetime import UTC, datetime
from email import policy
from email.parser import BytesParser
from io import BytesIO, StringIO
from pathlib import Path

SOURCE_DIR = Path("data/source")
PROCESSED_DIR = Path("data/raw")
SUPPORTED_SOURCE_EXTENSIONS = {
    ".txt",
    ".md",
    ".json",
    ".csv",
    ".pdf",
    ".docx",
    ".pptx",
    ".xlsx",
    ".eml",
    ".png",
    ".jpg",
    ".jpeg",
}
OCR_IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg"}
OCR_PDF_MIN_CHARS = 80


@dataclass(frozen=True)
class ExtractedSection:
    label: str
    text: str


@dataclass(frozen=True)
class ExtractionBundle:
    sections: list[ExtractedSection]
    method: str


@dataclass
class FileProcessingResult:
    source_path: Path
    output_path: Path | None
    status: str
    message: str


@dataclass
class IngestionSummary:
    processed: int = 0
    skipped: int = 0
    errors: int = 0
    processed_files: list[str] = field(default_factory=list)
    skipped_files: list[str] = field(default_factory=list)
    error_files: list[str] = field(default_factory=list)
    messages: list[str] = field(default_factory=list)


def ensure_data_directories() -> None:
    SOURCE_DIR.mkdir(parents=True, exist_ok=True)
    PROCESSED_DIR.mkdir(parents=True, exist_ok=True)


def sanitize_name(value: str) -> str:
    value = value.strip() or "document"
    return re.sub(r"[^A-Za-z0-9._-]+", "_", value)


def processed_filename_for(file_path: Path) -> str:
    try:
        relative_path = file_path.relative_to(SOURCE_DIR)
        pieces = list(relative_path.with_suffix("").parts)
    except ValueError:
        pieces = list(file_path.with_suffix("").parts)

    safe_base = "__".join(sanitize_name(piece) for piece in pieces if piece)
    extension_label = file_path.suffix.lower().lstrip(".") or "file"
    return f"{safe_base}__{extension_label}.txt"


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = text.replace("\x00", " ")
    text = html.unescape(text)
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


def _meaningful_char_count(text: str) -> int:
    return len(re.sub(r"[^A-Za-z0-9]+", "", text or ""))


def _clean_section_text(text: str) -> str:
    return normalize_text(text)


def _strip_html_tags(value: str) -> str:
    value = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", value)
    value = re.sub(r"(?s)<br\s*/?>", "\n", value)
    value = re.sub(r"(?s)</p>", "\n\n", value)
    value = re.sub(r"(?s)<[^>]+>", " ", value)
    return normalize_text(value)


def _build_bundle(section_map: list[tuple[str, str]], method: str) -> ExtractionBundle:
    sections = [
        ExtractedSection(label=label, text=_clean_section_text(text))
        for label, text in section_map
        if _clean_section_text(text)
    ]
    return ExtractionBundle(sections=sections, method=method)


def _extract_text_file(file_path: Path) -> ExtractionBundle:
    raw = file_path.read_bytes()
    return _build_bundle([("Body", raw.decode("utf-8", errors="replace"))], "utf-8 text decode")


def _extract_json_file(file_path: Path) -> ExtractionBundle:
    decoded = file_path.read_bytes().decode("utf-8", errors="replace")
    try:
        formatted = json.dumps(json.loads(decoded), indent=2, ensure_ascii=True)
    except json.JSONDecodeError:
        formatted = decoded
    return _build_bundle([("JSON", formatted)], "json decode")


def _extract_csv_file(file_path: Path) -> ExtractionBundle:
    decoded = file_path.read_bytes().decode("utf-8", errors="replace")
    reader = csv.reader(StringIO(decoded))
    lines = []
    for index, row in enumerate(reader, start=1):
        if not any(str(cell).strip() for cell in row):
            continue
        lines.append(f"Row {index}: " + " | ".join(str(cell) for cell in row))
    return _build_bundle([("CSV", "\n".join(lines))], "csv row extraction")


def _extract_docx_file(file_path: Path) -> ExtractionBundle:
    from docx import Document as DocxDocument

    document = DocxDocument(str(file_path))
    paragraphs = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]

    table_sections: list[str] = []
    for table_index, table in enumerate(document.tables, start=1):
        rows = []
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            table_sections.append(f"Table {table_index}\n" + "\n".join(rows))

    section_map: list[tuple[str, str]] = []
    if paragraphs:
        section_map.append(("Paragraphs", "\n".join(paragraphs)))
    section_map.extend((f"Table {i + 1}", text) for i, text in enumerate(table_sections))
    return _build_bundle(section_map, "docx paragraphs and tables")


def _extract_pptx_file(file_path: Path) -> ExtractionBundle:
    from pptx import Presentation

    presentation = Presentation(str(file_path))
    section_map: list[tuple[str, str]] = []

    for slide_index, slide in enumerate(presentation.slides, start=1):
        lines: list[str] = []
        title_shape = getattr(slide.shapes, "title", None)
        title_text = ""
        if title_shape is not None:
            title_text = _clean_section_text(getattr(title_shape, "text", ""))
            if title_text:
                lines.append(f"Slide title: {title_text}")

        for shape in slide.shapes:
            if shape is title_shape:
                continue
            if getattr(shape, "has_text_frame", False):
                text = _clean_section_text(shape.text)
                if text:
                    lines.append(text)
            if getattr(shape, "has_table", False):
                table_lines = []
                for row in shape.table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    if any(cells):
                        table_lines.append(" | ".join(cells))
                if table_lines:
                    lines.append("Table:\n" + "\n".join(table_lines))

        try:
            notes_slide = slide.notes_slide
            notes_frame = getattr(notes_slide, "notes_text_frame", None)
            notes_text = _clean_section_text(notes_frame.text if notes_frame is not None else "")
            if notes_text:
                lines.append("Speaker notes:\n" + notes_text)
        except Exception:
            pass

        if lines:
            label = f"Slide {slide_index}"
            if title_text:
                label = f"Slide {slide_index}: {title_text}"
            section_map.append((label, "\n\n".join(lines)))

    return _build_bundle(section_map, "pptx slide extraction")


def _extract_xlsx_file(file_path: Path) -> ExtractionBundle:
    from openpyxl import load_workbook

    workbook = load_workbook(filename=str(file_path), data_only=True, read_only=True)
    section_map: list[tuple[str, str]] = []

    for sheet in workbook.worksheets:
        lines = [f"Sheet name: {sheet.title}"]
        for row_index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
            values = ["" if value is None else str(value).strip() for value in row]
            if not any(values):
                continue
            lines.append(f"Row {row_index}: " + " | ".join(values))
        if len(lines) > 1:
            section_map.append((f"Sheet: {sheet.title}", "\n".join(lines)))

    workbook.close()
    return _build_bundle(section_map, "xlsx sheet extraction")


def _extract_eml_file(file_path: Path) -> ExtractionBundle:
    message = BytesParser(policy=policy.default).parsebytes(file_path.read_bytes())
    headers = [
        f"From: {message.get('from', '')}",
        f"To: {message.get('to', '')}",
        f"Cc: {message.get('cc', '')}",
        f"Bcc: {message.get('bcc', '')}",
        f"Subject: {message.get('subject', '')}",
        f"Date: {message.get('date', '')}",
    ]

    body_sections: list[str] = []
    if message.is_multipart():
        for part in message.walk():
            if part.get_content_disposition() == "attachment":
                continue
            content_type = part.get_content_type()
            try:
                content = part.get_content()
            except Exception:
                payload = part.get_payload(decode=True) or b""
                content = payload.decode("utf-8", errors="replace")

            if not isinstance(content, str):
                continue
            if content_type == "text/plain":
                cleaned = _clean_section_text(content)
            elif content_type == "text/html":
                cleaned = _strip_html_tags(content)
            else:
                continue
            if cleaned:
                body_sections.append(cleaned)
    else:
        content = message.get_content()
        if isinstance(content, str):
            body_sections.append(
                _strip_html_tags(content) if message.get_content_type() == "text/html" else _clean_section_text(content)
            )

    attachments = []
    for part in message.iter_attachments():
        filename = part.get_filename()
        if filename:
            attachments.append(filename)

    section_map = [
        ("Headers", "\n".join(line for line in headers if line.split(":", 1)[1].strip())),
        ("Body", "\n\n".join(section for section in body_sections if section)),
    ]
    if attachments:
        section_map.append(("Attachments", "\n".join(f"- {name}" for name in attachments)))
    return _build_bundle(section_map, "eml header/body extraction")


def _ocr_image(file_path: Path) -> str:
    from PIL import Image, ImageOps
    import pytesseract

    image = Image.open(file_path)
    image = ImageOps.exif_transpose(image).convert("L")
    return pytesseract.image_to_string(image)


def _extract_image_file(file_path: Path) -> ExtractionBundle:
    try:
        text = _ocr_image(file_path)
    except FileNotFoundError as exc:
        raise RuntimeError(
            "OCR failed because Tesseract is not installed or not on PATH."
        ) from exc
    return _build_bundle([("OCR Text", text)], "image OCR")


def _render_pdf_to_images(pdf_bytes: bytes) -> list[bytes]:
    import fitz

    document = fitz.open(stream=pdf_bytes, filetype="pdf")
    images: list[bytes] = []
    try:
        for page in document:
            pixmap = page.get_pixmap(matrix=fitz.Matrix(2, 2), alpha=False)
            images.append(pixmap.tobytes("png"))
    finally:
        document.close()
    return images


def _ocr_pdf_bytes(pdf_bytes: bytes) -> list[ExtractedSection]:
    from PIL import Image
    import pytesseract

    sections: list[ExtractedSection] = []
    for page_index, image_bytes in enumerate(_render_pdf_to_images(pdf_bytes), start=1):
        image = Image.open(BytesIO(image_bytes)).convert("L")
        text = _clean_section_text(pytesseract.image_to_string(image))
        if text:
            sections.append(ExtractedSection(label=f"Page {page_index} (OCR)", text=text))
    return sections


def _extract_pdf_file(file_path: Path) -> ExtractionBundle:
    from pypdf import PdfReader

    pdf_bytes = file_path.read_bytes()
    reader = PdfReader(BytesIO(pdf_bytes))
    sections: list[ExtractedSection] = []

    for page_index, page in enumerate(reader.pages, start=1):
        text = _clean_section_text(page.extract_text() or "")
        if text:
            sections.append(ExtractedSection(label=f"Page {page_index}", text=text))

    combined_text = "\n\n".join(section.text for section in sections)
    if _meaningful_char_count(combined_text) >= OCR_PDF_MIN_CHARS:
        return ExtractionBundle(sections=sections, method="pypdf text extraction")

    try:
        ocr_sections = _ocr_pdf_bytes(pdf_bytes)
    except FileNotFoundError as exc:
        raise RuntimeError(
            "PDF OCR fallback failed because Tesseract is not installed or not on PATH."
        ) from exc

    if ocr_sections:
        return ExtractionBundle(sections=ocr_sections, method="OCR fallback from rendered PDF pages")

    return ExtractionBundle(sections=sections, method="pypdf text extraction")


def extract_text_from_file(file_path: Path) -> str:
    return "\n\n".join(section.text for section in extract_sections_from_file(file_path).sections)


def extract_sections_from_file(file_path: Path) -> ExtractionBundle:
    suffix = file_path.suffix.lower()

    if suffix in {".txt", ".md"}:
        return _extract_text_file(file_path)
    if suffix == ".json":
        return _extract_json_file(file_path)
    if suffix == ".csv":
        return _extract_csv_file(file_path)
    if suffix == ".pdf":
        return _extract_pdf_file(file_path)
    if suffix == ".docx":
        return _extract_docx_file(file_path)
    if suffix == ".pptx":
        return _extract_pptx_file(file_path)
    if suffix == ".xlsx":
        return _extract_xlsx_file(file_path)
    if suffix == ".eml":
        return _extract_eml_file(file_path)
    if suffix in OCR_IMAGE_EXTENSIONS:
        return _extract_image_file(file_path)

    raise ValueError(f"Unsupported file type: {suffix}")


def _build_processed_output(source_path: Path, bundle: ExtractionBundle) -> str:
    try:
        relative_path = source_path.relative_to(SOURCE_DIR)
    except ValueError:
        relative_path = source_path

    header_lines = [
        f"Source filename: {source_path.name}",
        f"Source relative path: {relative_path.as_posix()}",
        f"Source type: {source_path.suffix.lower()}",
        f"Extraction method: {bundle.method}",
        f"Processed at: {datetime.now(UTC).isoformat()}",
        "",
    ]

    body_lines: list[str] = []
    for index, section in enumerate(bundle.sections, start=1):
        body_lines.extend(
            [
                f"## Section {index}: {section.label}",
                section.text,
                "",
            ]
        )

    return "\n".join(header_lines + body_lines).strip() + "\n"


def process_source_file(source_path: Path, processed_dir: Path) -> FileProcessingResult:
    suffix = source_path.suffix.lower()
    if suffix not in SUPPORTED_SOURCE_EXTENSIONS:
        return FileProcessingResult(
            source_path=source_path,
            output_path=None,
            status="skipped",
            message=f"[SKIPPED] {source_path.name} -> unsupported file type {suffix or '(none)'}",
        )

    try:
        bundle = extract_sections_from_file(source_path)
    except Exception as exc:
        return FileProcessingResult(
            source_path=source_path,
            output_path=None,
            status="error",
            message=f"[ERROR] {source_path.name} -> {exc}",
        )

    if not bundle.sections:
        return FileProcessingResult(
            source_path=source_path,
            output_path=None,
            status="skipped",
            message=f"[SKIPPED] {source_path.name} -> empty or unreadable after extraction",
        )

    processed_dir.mkdir(parents=True, exist_ok=True)
    output_path = processed_dir / processed_filename_for(source_path)

    try:
        output_text = _build_processed_output(source_path, bundle)
        output_path.write_text(output_text, encoding="utf-8")
    except Exception as exc:
        return FileProcessingResult(
            source_path=source_path,
            output_path=None,
            status="error",
            message=f"[ERROR] {source_path.name} -> failed to write processed output: {exc}",
        )

    return FileProcessingResult(
        source_path=source_path,
        output_path=output_path,
        status="processed",
        message=(
            f"[PROCESSED] {source_path.name} -> {output_path.name} "
            f"({len(bundle.sections)} sections, method: {bundle.method})"
        ),
    )


def ingest_source_directory(source_dir: Path, processed_dir: Path) -> IngestionSummary:
    summary = IngestionSummary()

    if not source_dir.exists():
        summary.messages.append(f"Source directory not found: {source_dir}")
        return summary

    source_files = sorted(path for path in source_dir.rglob("*") if path.is_file())
    if not source_files:
        summary.messages.append(f"No files found in source directory: {source_dir}")
        return summary

    for source_path in source_files:
        result = process_source_file(source_path, processed_dir)
        summary.messages.append(result.message)

        if result.status == "processed":
            summary.processed += 1
            summary.processed_files.append(result.output_path.name if result.output_path else source_path.name)
        elif result.status == "skipped":
            summary.skipped += 1
            summary.skipped_files.append(source_path.name)
        else:
            summary.errors += 1
            summary.error_files.append(source_path.name)

    return summary


def count_supported_source_files(source_dir: Path) -> int:
    if not source_dir.exists():
        return 0
    return sum(
        1
        for path in source_dir.rglob("*")
        if path.is_file() and path.suffix.lower() in SUPPORTED_SOURCE_EXTENSIONS
    )


def get_processed_file_count(processed_dir: Path) -> int:
    if not processed_dir.exists():
        return 0
    return sum(1 for path in processed_dir.rglob("*.txt") if path.is_file())
